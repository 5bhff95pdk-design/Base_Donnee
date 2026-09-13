"""Relations explicites entre personnages : validation et exports sans inférence.

La source éditoriale est data/relations.csv. Les liens ne sont jamais recalculés
à partir des noms, des clans ou de l'analyse automatique du texte.
"""
import csv
import json
from pathlib import Path

SOURCE = Path('data/relations.csv')
EXPORT = Path('relations_personnages.json')
COLONNES = ['Source_ID', 'Type', 'Cible_ID', 'Preuve_ID', 'Champ_source', 'Extrait']
# Source → cible ; les relations symétriques sont stockées une seule fois.
TYPES = {
    'parent_de': ('enfant_de', False),
    'conjoint_de': ('conjoint_de', True),
    'ex_conjoint_de': ('ex_conjoint_de', True),
    'fratrie': ('fratrie', True),
    'cousin_de': ('cousin_de', True),
    'collegue_de': ('collegue_de', True),
    'colocataire_de': ('colocataire_de', True),
    'superieur_de': ('subordonne_de', False),
    'locataire_de': ('bailleur_de', False),
    'enseignant_de': ('eleve_de', False),
    # Liens de famille que le champ Parenté énonce explicitement sans les
    # réduire à une filiation : oncle/tante, grand-parent, parrain/marraine.
    # Sans ces types, ces liens restaient en prose et absents de la table.
    'oncle_de': ('neveu_de', False),
    'grand_parent_de': ('petit_enfant_de', False),
    'parrain_de': ('filleul_de', False),
}


def valider(lignes, personnages):
    """Vérifie références, unicité, provenance et absence de cycles parentaux."""
    base = {r['ID']: r for r in personnages}
    vus = set()
    parents = {}
    for numero, r in enumerate(lignes, 2):
        def erreur(message):
            raise SystemExit(f'✘ Relations ligne {numero} : {message}')

        if set(r) != set(COLONNES) or any(not isinstance(v, str) or not v.strip()
                                         for v in r.values()):
            erreur('colonnes ou valeurs manquantes')
        for champ in ('Source_ID', 'Cible_ID', 'Preuve_ID'):
            if r[champ] not in base:
                erreur(f'{champ} inconnu : {r[champ]}')
        a, b, type_lien = r['Source_ID'], r['Cible_ID'], r['Type']
        if a == b:
            erreur('auto-relation interdite')
        if type_lien not in TYPES:
            erreur(f'type inconnu : {type_lien}')
        if TYPES[type_lien][1] and a > b:
            erreur('relation symétrique : ranger les deux IDs par ordre lexical')
        cle = (a, type_lien, b)
        if cle in vus:
            erreur('relation en double')
        vus.add(cle)
        # Première tranche : seuls les liens décrits dans Parenté sont admis.
        if r['Champ_source'] != 'Parenté':
            erreur('champ de provenance non pris en charge')
        if r['Preuve_ID'] not in (a, b):
            erreur('la preuve doit appartenir à une extrémité du lien')
        if r['Extrait'] != base[r['Preuve_ID']]['Parenté']:
            erreur('preuve désynchronisée : relire la relation avant de mettre à jour l’extrait')
        if type_lien == 'parent_de':
            parents.setdefault(a, []).append(b)

    # Un parent ne peut pas être son propre descendant. Pas d'inférence par âge :
    # les données ne distinguent pas encore filiation biologique/adoptive.
    actifs, termines = set(), set()

    def visiter(identifiant):
        if identifiant in actifs:
            raise SystemExit('✘ Relations : cycle de parenté')
        if identifiant in termines:
            return
        actifs.add(identifiant)
        for enfant in parents.get(identifiant, []):
            visiter(enfant)
        actifs.remove(identifiant)
        termines.add(identifiant)

    for identifiant in parents:
        visiter(identifiant)
    return sorted(lignes, key=lambda r: (r['Source_ID'], r['Type'], r['Cible_ID']))


def charger(personnages):
    if not SOURCE.is_file():
        raise SystemExit(f'✘ Source des relations introuvable : {SOURCE}')
    with SOURCE.open(encoding='utf-8', newline='') as f:
        lecteur = csv.DictReader(f, delimiter=';')
        if lecteur.fieldnames != COLONNES:
            raise SystemExit('✘ Relations : en-tête invalide')
        lignes = list(lecteur)
    if not lignes:
        raise SystemExit('✘ Relations : source vide')
    return valider(lignes, personnages)


def serialiser(lignes, personnages):
    noms = {r['ID']: r['Nom'] for r in personnages}
    return {
        'version_schema': 1,
        'perimetre': 'Extraction éditoriale partielle des liens explicites de Parenté ; sans inférence.',
        'types': {k: {'inverse': v[0], 'symetrique': v[1]} for k, v in TYPES.items()},
        'relations': [{
            'source_id': r['Source_ID'], 'source_nom': noms[r['Source_ID']],
            'type': r['Type'],
            'cible_id': r['Cible_ID'], 'cible_nom': noms[r['Cible_ID']],
            'preuve': {'fichier': 'data/personnages.csv', 'personnage_id': r['Preuve_ID'],
                       'champ': r['Champ_source'], 'extrait': r['Extrait']},
        } for r in lignes],
    }


def exporter(lignes, personnages):
    EXPORT.write_text(json.dumps(serialiser(lignes, personnages), ensure_ascii=False, indent=2)
                      + '\n', encoding='utf-8')


def ecrire_feuille(classeur, lignes, personnages):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    noms = {r['ID']: r['Nom'] for r in personnages}
    feuille = classeur.create_sheet('Relations')
    entetes = ['Source_ID', 'Source_nom', 'Type', 'Cible_ID', 'Cible_nom',
               'Preuve_ID', 'Champ_source', 'Extrait']
    feuille.append(entetes)
    for r in lignes:
        feuille.append([r['Source_ID'], noms[r['Source_ID']], r['Type'], r['Cible_ID'],
                        noms[r['Cible_ID']], r['Preuve_ID'], r['Champ_source'], r['Extrait']])
    for cell in feuille[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1F3864')
    for i, largeur in enumerate([12, 26, 22, 12, 26, 12, 18, 75], 1):
        feuille.column_dimensions[get_column_letter(i)].width = largeur
    for row in feuille.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    feuille.freeze_panes = 'C2'
    feuille.auto_filter.ref = feuille.dimensions
