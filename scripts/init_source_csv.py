#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
scripts/init_source_csv.py — migration UNIQUE du classeur XLSX vers la source
de vérité texte : data/personnages.csv (+ data/narration.csv).

Historique : jusqu'au 2026-09-13, la source de vérité était
base_personnages_fictifs.xlsx, un classeur binaire que construire_base.py
réécrivait lui-même (aller-retour openpyxl) — diff Git illisible, fusion
impossible en cas de conflit, perte des commentaires et des mises en forme.
Ce script bascule la source en CSV UTF-8 : lisible, diffable, éditable.

Produit :
  data/personnages.csv   16 colonnes (source publique, versionnée)
  data/narration.csv     8 colonnes  (source privée, NON versionnée → .gitignore)

Transformations appliquées (une seule fois) :
  1. Famille « Lavoie (JP) » → Famille « Lavoie » + Branche « JP »
     (la convention du qualificatif entre parenthèses n'était documentée
     nulle part ; elle devient une colonne à part entière).
  2. Faction : 128 valeurs libres → 17 factions canoniques
     (data/factions.txt) ; l'original est conservé dans « Faction (détail) ».

Idempotent : refuse d'écraser une source existante sans --force.
Usage :  python3 scripts/init_source_csv.py [--force]
"""
import argparse
import csv
import os
import re
import sys

import openpyxl

RACINE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(RACINE, 'base_personnages_fictifs.xlsx')
DATA = os.path.join(RACINE, 'data')
PERSONNAGES_CSV = os.path.join(DATA, 'personnages.csv')
NARRATION_CSV = os.path.join(DATA, 'narration.csv')
FACTIONS_TXT = os.path.join(DATA, 'factions.txt')

COLONNES = ['Nom', 'Surnom', 'Type', 'Age', 'Rôle', 'Secteur', 'Adresse',
            'Latitude', 'Longitude', 'Apparence', 'Vêtements', 'Tic / Objet',
            'Portrait', 'Famille', 'Branche', 'Parenté']
NARR_COLS = ['Nom', 'Surnom', 'Faction', 'Faction (détail)', 'Lien Spot',
             'Quote joual', 'Arc S1', 'Age apparent']

# ---------------------------------------------------------------- factions
# 128 valeurs libres → 17 factions canoniques. Table lue à l'envers : c'est la
# liste de droite qui fait foi, le script échoue s'il reste une valeur non
# classée (protection contre l'oubli d'une valeur ajoutée depuis).
FACTIONS = {
    'Ruelle': [
        'Ruelle', "Ruelle d'en Face", "Ruelle d'en Face / Famille",
        'Ruelle / Bérubé', 'Ruelle / Gratien', 'Ruelle / Lavoie',
        'Ruelle / Louche', 'Ruelle / Marge', 'Ruelle / Neutre',
        'Ruelle / Salon', 'Ruelle / Survivaliste', 'Réseau / Louche',
        'Quartier / Marge',
    ],
    'Sous-sol': ['Sous-sol'],
    'Bloc': [
        '1er étage', '2e étage / Louche numérique', '3e étage / Louche numérique',
        "Bloc d'en Face", "Bloc d'en Face / Famille", "Bloc d'en face",
        "Bloc d'à côté / Louche sage", 'Locataires / Neutre', 'Vieux bloc',
        'Vieux bloc / Port', 'Westmount / Louche naïf',
    ],
    # « Aide / Famille » : la posture « Aide » porte sur le foyer.
    'Famille': [
        'Famille / Alex', 'Famille / Aubin', 'Famille / Bérubé',
        'Famille / Chicoine', 'Famille / Chloé', 'Famille / Gratien',
        'Famille / Ruelle', 'Famille / Truchon', 'Famille JP',
        'Famille JP / Aréna', 'Famille Marco', 'Aide / Famille',
    ],
    'Commerce': [
        'Argent', 'Argent / Chambre', 'Argent / Famille', 'Argent / Ruelle',
        'Au-dessus Dépanneur / Louche', 'Casse', 'Commerce / Neutre',
        'Commerces / Famille', 'Dépanneur / Famille', 'Dépanneur / Neutre',
        'Garage / Neutre', 'Rue / Louche entrepreneure',
        'Station-service / Famille', 'Taxi / Neutre',
    ],
    'Bar du Coin': ['Bar du Coin', 'Bar du Coin / Famille', 'Bar du Coin / Louche'],
    'Institutions': [
        'Aréna', 'Aréna / Neutre', 'Armée / Bagotville', 'Bagotville / Marge',
        'Bibliothèque / Neutre', 'Institutions', 'Politique', 'Poste / Neutre',
        'STS / Neutre', 'Tourisme / Neutre', 'Ville', 'Ville / Argent',
        'Ville / Famille', 'Ville / Neutre', 'Ville / Tourisme',
    ],
    'Santé / Enseignement': [
        'Aide', 'Aide / CHSLD', 'Aide / Université', 'CLSC / Aide', 'CPE / Aide',
        'Caserne / Urgences', 'Cégep / Naïf', 'Hôpital', 'Hôpital / Aide',
        'Urgences / Aide', 'École', 'École / Aide', 'ATM / École',
    ],
    'Église': ['Église', 'Église / Diocèse', 'Église / Marge'],
    'Santini': [
        'Santini', 'Santini / Argent', 'Santini / Conseiller', 'Santini / Famille',
        'Santini / Façade', 'Santini / Port', 'Santini / Quai', 'Santini / Relève',
        'Santini / Ruelle', 'Santini / Sommet',
    ],
    'Motards': ['Motards', 'Motards / Aînée', 'Motards / Louche', 'Motards / Louves'],
    'Industrie': [
        'Aluminerie / Direction', 'Aluminerie / Syndicat', 'Aluminerie / Syndiqué',
        'Hydro / Syndicat',
    ],
    'Média / Art': [
        'ATM / Art', 'Art', 'Art / Musique', 'Culture / Neutre', 'Média',
        'Média / Rang', 'Presse', 'Radio pirate / Marge',
    ],
    'Plein air / Port': [
        'Bec-Scie / Neutre', 'Campagne', 'Campagne / Neutre', 'Fjord / Neutre',
        'Parc du Fjord', 'Port', 'Port / Neutre', 'Pyramide / Neutre',
    ],
    'Loi / Sécurité': [
        'Loi', 'Loi / Direction', 'Loi / Famille', 'Loi / Magistrature',
        'Sécurité / Louche ex-loi',
    ],
    'Propre': ['Propre', 'Santé / Propre'],
    'Neutre': ['Neutre', 'Rue / Neutre'],
}


def factions_canoniques():
    """Lit data/factions.txt (source de vérité du vocabulaire)."""
    if not os.path.exists(FACTIONS_TXT):
        sys.exit(f"✘ {FACTIONS_TXT} introuvable")
    with open(FACTIONS_TXT, encoding='utf-8') as f:
        return [lg.strip() for lg in f
                if lg.strip() and not lg.lstrip().startswith('#')]


def table_factions():
    """Inverse FACTIONS après avoir vérifié qu'elle colle à data/factions.txt."""
    canoniques = factions_canoniques()
    inconnues = [c for c in FACTIONS if c not in canoniques]
    if inconnues:
        sys.exit(f"✘ Faction(s) absente(s) de {FACTIONS_TXT} : {inconnues}")
    table = {}
    for canonique, sources in FACTIONS.items():
        for src in sources:
            if src in table and table[src] != canonique:
                sys.exit(f"✘ Valeur « {src} » classée dans deux factions")
            table[src] = canonique
    return table, canoniques


def separer_famille(valeur):
    """"Lavoie (JP)" → ("Lavoie", "JP") ; "Santini" → ("Santini", "")."""
    v = (valeur or '').strip()
    m = re.match(r'^(.*?)\s*\(([^()]*)\)\s*$', v)
    if m:
        return m.group(1).strip(), m.group(2).strip()
    return v, ''


def nombre(v):
    """Rend le nombre tel qu'il sera relu sans perte (entier nu si possible)."""
    if v is None or v == '':
        return ''
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return repr(float(v)) if isinstance(v, float) else str(v)


def lire_xlsx():
    if not os.path.exists(XLSX):
        sys.exit(f"✘ {XLSX} introuvable (déjà migré ?)")
    wb = openpyxl.load_workbook(XLSX)
    ws = wb['Personnages']
    hdr = [c.value for c in ws[1]]
    persos = [dict(zip(hdr, [c.value for c in r]))
              for r in ws.iter_rows(min_row=2)
              if any(c.value is not None for c in r)]
    narr = []
    if 'Narration' in wb.sheetnames:
        nw = wb['Narration']
        nh = [c.value for c in nw[1]]
        narr = [dict(zip(nh, [c.value for c in r]))
                for r in nw.iter_rows(min_row=2)
                if any(c.value is not None for c in r)]
    return persos, narr


def ecrire_personnages(persos):
    with open(PERSONNAGES_CSV, 'w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, fieldnames=COLONNES, delimiter=';',
                           lineterminator='\n')
        w.writeheader()
        for p in persos:
            famille, branche = separer_famille(p.get('Famille'))
            ligne = {c: ('' if p.get(c) is None else p.get(c)) for c in COLONNES}
            ligne['Famille'] = famille
            ligne['Branche'] = branche
            for c in ('Age', 'Latitude', 'Longitude'):
                ligne[c] = nombre(ligne[c])
            w.writerow(ligne)


def ecrire_narration(narr, table):
    with open(NARRATION_CSV, 'w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, fieldnames=NARR_COLS, delimiter=';',
                           lineterminator='\n')
        w.writeheader()
        for n in narr:
            source = (n.get('Faction') or '').strip()
            if source not in table:
                sys.exit(f"✘ Faction non classée : {source!r} ({n.get('Nom')})")
            w.writerow({'Nom': n.get('Nom') or '',
                        'Surnom': n.get('Surnom') or '',
                        'Faction': table[source],
                        'Faction (détail)': source,
                        'Lien Spot': n.get('Lien Spot') or '',
                        'Quote joual': n.get('Quote joual') or '',
                        'Arc S1': n.get('Arc S1') or '',
                        'Age apparent': nombre(n.get('Age apparent'))})


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--force', action='store_true',
                    help='écrase data/*.csv s’ils existent déjà')
    args = ap.parse_args()

    os.makedirs(DATA, exist_ok=True)
    for cible in (PERSONNAGES_CSV, NARRATION_CSV):
        if os.path.exists(cible) and not args.force:
            sys.exit(f"✘ {cible} existe déjà (utiliser --force pour écraser)")

    table, canoniques = table_factions()

    persos, narr = lire_xlsx()
    ecrire_personnages(persos)
    print(f"  ✔ data/personnages.csv — {len(persos)} entrées × {len(COLONNES)} colonnes")
    if narr:
        ecrire_narration(narr, table)
        print(f"  ✔ data/narration.csv — {len(narr)} lignes, "
              f"factions normalisées ({len(canoniques)} valeurs canoniques)")
    print("\nSource texte en place. Vérifier puis relancer : python3 construire_base.py")


if __name__ == '__main__':
    main()
