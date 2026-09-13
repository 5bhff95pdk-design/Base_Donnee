# -*- coding: utf-8 -*-
"""
tests/test_base.py — garde-fous de la base de personnages fictifs.

Lancé par la CI (.github/workflows/validation.yml) après régénération :
    python -m unittest discover -s tests -v

Couvre les 13 points de la révision du 2026-09-12 : effectifs annoncés,
conventions d'écriture (adresses, rôles, surnoms), cohérence géographique,
couverture de la narration, angle mort démographique, portraits et planche
contact, carte hors ligne, licences distinctes et reproductibilité des
livrables (vérifiée dans la CI par « git diff » après régénération).
"""
import csv
import glob
import json
import os
import re
import unittest
from pathlib import Path

import openpyxl

RACINE = Path(__file__).resolve().parent.parent
os.chdir(RACINE)

N = 173                      # nombre canonique d'entrées
COLONNES = ['Nom', 'Surnom', 'Type', 'Age', 'Rôle', 'Secteur', 'Adresse',
            'Latitude', 'Longitude', 'Apparence', 'Vêtements',
            'Tic / Objet', 'Portrait', 'Famille', 'Parenté']
MINEURS = {'Léo Cloutier': 9, 'Nour Benali': 13, 'Alexandre Lavoie': 16}

# Boîtes approximatives par arrondissement (validées sur les rues réelles).
BOITES = {
    'La Baie':     (48.28, 48.381, -71.02, -70.75),
    'Chicoutimi':  (48.38, 48.45,  -71.12, -71.00),
    'Jonquière':   (48.40, 48.45,  -71.30, -71.15),
}


def charger_xlsx():
    wb = openpyxl.load_workbook('base_personnages_fictifs.xlsx')
    ws = wb['Personnages']
    hdr = [c.value for c in ws[1]]
    recs = [dict(zip(hdr, [c.value for c in r]))
            for r in ws.iter_rows(min_row=2)
            if any(c.value is not None for c in r)]
    return wb, recs


class TestEffectifsEtFormats(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.wb, cls.recs = charger_xlsx()

    def test_nombre_entrees(self):
        self.assertEqual(len(self.recs), N, 'le classeur maître doit contenir 173 entrées')

    def test_les_quatre_exports_ont_le_meme_effectif(self):
        with open('base_personnages_fictifs.json', encoding='utf-8') as f:
            js = json.load(f)
        with open('base_personnages_fictifs.geojson', encoding='utf-8') as f:
            gj = json.load(f)
        with open('base_personnages_fictifs.csv', encoding='utf-8-sig') as f:
            lignes = list(csv.DictReader(f, delimiter=';'))
        self.assertEqual(len(js), N)
        self.assertEqual(len(gj['features']), N)
        self.assertEqual(len(lignes), N)

    def test_colonnes_canoniques(self):
        hdr = [c.value for c in self.wb['Personnages'][1]]
        self.assertEqual(hdr, COLONNES)


class TestConventions(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.wb, cls.recs = charger_xlsx()

    def test_noms_uniques(self):
        noms = [r['Nom'] for r in self.recs]
        doublons = {n for n in noms if noms.count(n) > 1}
        self.assertFalse(doublons, f'noms en double : {doublons}')

    def test_pas_de_surnom_incruste_dans_le_nom(self):
        # #6 : le surnom vit dans la colonne Surnom, jamais entre « » dans Nom.
        for r in self.recs:
            self.assertNotIn('«', r['Nom'], f'{r["Nom"]} : surnom dans Nom')
            self.assertNotIn('»', r['Nom'], f'{r["Nom"]} : surnom dans Nom')
        bernard = next(r for r in self.recs if r['Nom'] == 'Bernard Tremblay')
        karine = next(r for r in self.recs if r['Nom'] == 'Karine Desrosiers')
        self.assertEqual(bernard['Surnom'], 'Bunker')
        self.assertEqual(karine['Surnom'], 'Zoom')

    def test_adresses_avec_numero_ou_lieu_dit(self):
        # #6 : soit « Numéro, Rue », soit le préfixe explicite « Lieu-dit : ».
        for r in self.recs:
            ad = str(r['Adresse']).strip()
            avec_num = re.match(r'^\d+[\s,]', ad)
            lieu_dit = ad.startswith('Lieu-dit :')
            self.assertTrue(avec_num or lieu_dit,
                            f'{r["Nom"]} : adresse sans numéro ni « Lieu-dit : » → {ad!r}')
        # les 7 cas non adressables portent bien le préfixe
        pour = ['Bernard Tremblay', 'Félix Otis', 'Gaston Ha! Ha!', 'Justine Fortin',
                'Luc Villeneuve', 'Natasha Vollant', 'William McKenzie']
        for nom in pour:
            self.assertTrue(next(r for r in self.recs if r['Nom'] == nom)['Adresse']
                            .startswith('Lieu-dit :'))

    def test_roles_sans_clan(self):
        # #6 : ni « clan », ni nom de famille/faction dans le Rôle.
        for r in self.recs:
            role = str(r['Rôle'])
            self.assertNotRegex(role.lower(), r'\bclan\b',
                                f'{r["Nom"]} : {role}')
            self.assertNotIn('Santini', role, f'{r["Nom"]} : Rôle cité la famille → {role}')
        # la famille reste portée par la colonne Famille
        vittorio = next(r for r in self.recs if r['Nom'] == 'Vittorio Santini')
        self.assertEqual(vittorio['Famille'], 'Santini')

    def test_pas_la_coquille_college(self):
        # #5 : « collège de Laurie Lavoie » → « collègue de Laurie Lavoie ».
        for r in self.recs:
            for champ in ('Parenté', 'Rôle'):
                v = r.get(champ)
                if isinstance(v, str):
                    self.assertNotIn('collège de', v, f'{r["Nom"]} : coquille « collège »')
        sophie = next(r for r in self.recs if r['Nom'] == 'Sophie Labbé')
        self.assertIn('collègue de Laurie Lavoie', sophie['Parenté'])


class TestGeographie(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        _, cls.recs = charger_xlsx()

    def test_coordonnees_dans_le_bon_arrondissement(self):
        # #4 : Danny Fortin était sectorisé Chicoutimi mais géolocalisé à
        # Jonquière (2140, boulevard Harvey). Toutes les coords doivent
        # tomber dans la boîte de leur secteur.
        for r in self.recs:
            lat, lon = float(r['Latitude']), float(r['Longitude'])
            amin, amax, omin, omax = BOITES[r['Secteur']]
            self.assertTrue(amin <= lat <= amax and omin <= lon <= omax,
                            f'{r["Nom"]} ({r["Secteur"]}) hors boîte : {lat},{lon}')

    def test_danny_fortin_est_bien_a_chicoutimi(self):
        d = next(r for r in self.recs if r['Nom'] == 'Danny Fortin')
        self.assertEqual(d['Secteur'], 'Chicoutimi')
        self.assertLess(float(d['Longitude']), -71.0)
        self.assertNotIn('Harvey', str(d['Adresse']))


class TestDemographie(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        _, cls.recs = charger_xlsx()

    def test_des_mineurs_humains_existent(self):
        # #10 : avant correction, le seul < 18 ans était le chat (7 ans).
        humains = [r for r in self.recs if r['Type'] == 'Humain']
        mineurs = [r for r in humains if int(r['Age']) < 18]
        self.assertEqual({r['Nom'] for r in mineurs}, set(MINEURS))
        for nom, age in MINEURS.items():
            self.assertEqual(int(next(r for r in self.recs if r['Nom'] == nom)['Age']), age)
        ages = sorted(int(r['Age']) for r in humains)
        self.assertEqual(ages[0], 9)

    def test_un_seul_animal(self):
        animaux = [r for r in self.recs if r['Type'] == 'Animal']
        self.assertEqual([r['Nom'] for r in animaux], ['Pisse-Feu'])


class TestNarration(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.wb, cls.recs = charger_xlsx()

    def test_couverture_100_pour_cent(self):
        # #9 : 40/170 → toutes les entrées doivent avoir une ligne complète.
        ws = self.wb['Narration']
        hdr = [c.value for c in ws[1]]
        lignes = [dict(zip(hdr, [c.value for c in r]))
                  for r in ws.iter_rows(min_row=2)
                  if any(c.value is not None for c in r)]
        noms_base = {r['Nom'] for r in self.recs}
        noms_narr = {l['Nom'] for l in lignes}
        self.assertEqual(noms_narr, noms_base,
                         f'manquants : {noms_base - noms_narr} ; orphelins : {noms_narr - noms_base}')
        for l in lignes:
            for champ in ('Faction', 'Lien Spot', 'Quote joual', 'Arc S1'):
                self.assertTrue(str(l.get(champ) or '').strip(),
                                f'{l["Nom"]} : champ Narration vide « {champ} »')

    def test_narration_pas_exportee(self):
        with open('base_personnages_fictifs.geojson', encoding='utf-8') as fh:
            gj = json.load(fh)
        for feat in gj['features']:
            self.assertNotIn('faction', feat['properties'])
            self.assertNotIn('quote', str(feat['properties']).lower())


class TestPortraits(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        _, cls.recs = charger_xlsx()

    def test_deux_gabarits_pour_chaque_personnage(self):
        for r in self.recs:
            web = r['Portrait']
            self.assertTrue(web and os.path.exists(web),
                            f'{r["Nom"]} : portrait web manquant ({web})')
            vig = web.replace('-web.webp', '-vignette.webp')
            self.assertTrue(os.path.exists(vig), f'{r["Nom"]} : vignette manquante ({vig})')

    def test_pas_de_gabarit_archive(self):
        # #8 : le gabarit .webp sans suffixe (~140 Ko) n'est référencé par rien.
        archives = [f for f in glob.glob('portraits/*.webp')
                    if not f.endswith('-web.webp')
                    and not f.endswith('-vignette.webp')
                    and not f.endswith('planche-contact-generale.webp')]
        self.assertEqual(archives, [], f'gabarits archive non référencés : {archives}')

    def test_planche_contact_generee(self):
        # #2 : planche-contact-generale.webp, livrable annoncé et désormais présent.
        self.assertTrue(os.path.exists('portraits/planche-contact-generale.webp'))
        try:
            from PIL import Image
            with Image.open('portraits/planche-contact-generale.webp') as im:
                self.assertEqual(im.width, 2510)  # 10 colonnes × 240 px + marges
        except ImportError:
            self.skipTest('Pillow absent')

    def test_nouveaux_portraits_numerotes(self):
        for base in ('171-leo-cloutier', '172-nour-benali', '173-alexandre-lavoie'):
            self.assertTrue(os.path.exists(f'portraits/{base}-web.webp'), base)
            self.assertTrue(os.path.exists(f'portraits/{base}-vignette.webp'), base)


class TestCarte(unittest.TestCase):
    def test_leaflet_vendore_en_local(self):
        # #13 : plus de CDN unpkg ; le CODE de la carte est utilisable hors
        # ligne. Les tuiles RESTENT des services en ligne (voir le README).
        self.assertTrue(os.path.exists('carte/vendor/leaflet/leaflet.js'))
        self.assertTrue(os.path.exists('carte/vendor/leaflet/leaflet.css'))
        self.assertTrue(os.path.exists('carte/vendor/leaflet/LICENSE'))
        cdn_code = re.compile(r'https?://[^"\' ]*(unpkg|jsdelivr|cdnjs|googleapis)', re.I)
        for f in ('carte-la-baie-saguenay.html', 'carte/index.html'):
            h = open(f, encoding='utf-8').read()
            self.assertIsNone(cdn_code.search(h), f'{f} : dépendance CDN pour le code')
        with open('carte-la-baie-saguenay.html', encoding='utf-8') as fh:
            racine = fh.read()
        self.assertIn('carte/vendor/leaflet/leaflet.js', racine)
        with open('carte/index.html', encoding='utf-8') as fh:
            index = fh.read()
        self.assertIn('"vendor/leaflet/leaflet.js"', index.replace("'", '"'))

    def test_persos_injectes_dans_les_deux_cartes(self):
        for f in ('carte-la-baie-saguenay.html', 'carte/index.html'):
            with open(f, encoding='utf-8') as fh:
                h = fh.read()
            m = re.search(r'const PERSOS=(\[.*?\]);', h, flags=re.S)
            self.assertIsNotNone(m, f'{f} : PERSOS introuvable')
            self.assertEqual(len(json.loads(m.group(1))), N, f'{f} : effectif PERSOS')

    def test_licence_leaflet_presente(self):
        with open('carte/vendor/leaflet/LICENSE', encoding='utf-8') as fh:
            txt = fh.read()
        self.assertIn('BSD', txt)


class TestDocumentationEtLicences(unittest.TestCase):
    def test_readme_md_est_la_vraie_documentation(self):
        # #3 : la doc doit être visible sur la page GitHub (README.md), pas
        # planquée dans README.txt.
        self.assertTrue(os.path.exists('README.md'))
        self.assertFalse(os.path.exists('README.txt'),
                         'README.txt fait doublon avec README.md')
        with open('README.md', encoding='utf-8') as fh:
            h = fh.read()
        for mot in ('173', 'ODbL', 'MIT', 'générées par IA',
                    'planche-contact-generale', 'Lieu-dit'):
            self.assertIn(mot, h, f'README.md oublie « {mot} »')

    def test_trois_statuts_de_licence_distincts(self):
        # #12 : code MIT, données dérivées OSM en ODbL, portraits IA séparés.
        self.assertTrue(os.path.exists('LICENSE-DONNEES.md'))
        with open('LICENSE-DONNEES.md', encoding='utf-8') as fh:
            d = fh.read()
        for mot in ('MIT', 'ODbL', 'OpenStreetMap', 'CC BY 4.0', 'IA'):
            self.assertIn(mot, d, f'LICENSE-DONNEES.md oublie « {mot} »')
        with open('LICENSE', encoding='utf-8') as fh:
            mit = fh.read()
        self.assertIn('MIT License', mit)

    def test_lisez_moi_annonce_le_bon_effectif(self):
        # #1 : le classeur ne doit plus annoncer 140 personnages / 160 lignes.
        wb, _ = charger_xlsx()
        textes = ' '.join(str(c.value) for row in wb['Lisez-moi'].iter_rows() for c in row
                          if c.value)
        self.assertIn(f'{N} entrées', textes)
        self.assertNotIn('140 personnages', textes)
        self.assertNotIn('160 lignes', textes)

    def test_outillage_present(self):
        # #11 : dépendances, ignores, attributs, tests et CI.
        for f in ('requirements.txt', '.gitignore', '.gitattributes',
                  '.github/workflows/validation.yml',
                  'scripts/retirer_archives_git.sh'):
            self.assertTrue(os.path.exists(f), f'{f} absent')


class TestReproductibilite(unittest.TestCase):
    def test_pas_d_horloge_dans_le_xlsx(self):
        # #7 : « Généré le… » est figé ; pas de date du jour dans le classeur.
        import zipfile
        with zipfile.ZipFile('base_personnages_fictifs.xlsx') as z:
            core = z.read('docProps/core.xml').decode('utf-8')
        self.assertIn('2026-09-11T00:00:00Z', core)
        created = re.search(r'<dcterms:created[^>]*>([^<]+)</dcterms:created>', core).group(1)
        modified = re.search(r'<dcterms:modified[^>]*>([^<]+)</dcterms:modified>', core).group(1)
        self.assertEqual(created, modified)


if __name__ == '__main__':
    unittest.main(verbosity=2)
